#!/usr/bin/env python3
"""Aggregate scene-level and object-level statistics from evaluator metrics files.

Reads an evaluation manifest (a JSON list, or an object with an ``items``,
``data``, or ``samples`` list) in which each record points to a results
directory containing ``metrics*.json`` written by the paper evaluator. The
script aggregates per-scene and per-object CD/IoU across all records and prints
or writes the statistics. With ``--out-excel`` the workbook also contains
per-sample rows and an optional combined per-sample sheet across the
SAM3D/ours/MIDI/PartCrafter result keys.

Example:
  python scripts/compute_stats.py \\
    --config outputs/eval.json \\
    --results-key sam3d_results_path \\
    --out-json outputs/metrics_stats.json
"""

from __future__ import annotations

import argparse
import json
import statistics
import sys
import re
from pathlib import Path
from typing import Any, Dict, List, Optional


_METRIC_KEYS = [
    "scene_cd",
    "scene_iou",
    "object_cd",
    "object_iou",
]
_DEFAULT_COMBINED_RESULTS_KEYS = (
    "sam3d_results_path",
    "ours_results_path",
    "midi_results_path",
    "partcrafter_results_path",
)
_COMBINED_GLOB_OVERRIDES = {
    "midi_results_path": "metrics*",
    "partcrafter_results_path": "metrics*",
}


def _sub_scene_id(rec: Dict[str, Any]) -> Any:
    if "sub_scene_id" in rec:
        return rec.get("sub_scene_id")
    return rec.get("subscene_id")


_DIFFICULTY_KEYS = ("easy", "medium", "hard", "unknown")


def _difficulty_from_rec(rec: Dict[str, Any]) -> str:
    """
    Infer difficulty (easy/medium/hard) from sub_scene_id patterns like:
      images-s1e1-png, images-s1m1-png, images-s1h1-png
    Fallback to 'unknown' when the pattern is not found.
    """
    sid = str(_sub_scene_id(rec) or rec.get("image_path") or "")
    # Look for s<digits><e|m|h><digits>, e.g. s1e1 / s10m2 / s3h1
    m = re.search(r"s\d+([emh])\d", sid)
    if not m:
        return "unknown"
    ch = m.group(1)
    if ch == "e":
        return "easy"
    if ch == "m":
        return "medium"
    if ch == "h":
        return "hard"
    return "unknown"


def _prefix_for_results_key(results_key: str) -> str:
    suffix = "_results_path"
    if results_key.endswith(suffix):
        return results_key[: -len(suffix)]
    return results_key


def _combined_metrics_glob(results_key: str, default_glob: str) -> str:
    return _COMBINED_GLOB_OVERRIDES.get(results_key, default_glob)


def _safe_float(x: Any) -> Optional[float]:
    try:
        if x is None:
            return None
        v = float(x)
        if v != v:  # NaN
            return None
        return v
    except Exception:
        return None


def _pick_metrics_file(dir_path: Path, pattern: str, prefer_latest: bool) -> Optional[Path]:
    files = [p for p in dir_path.glob(pattern) if p.is_file()]
    if not files:
        return None
    if len(files) == 1:
        return files[0]
    if prefer_latest:
        files.sort(key=lambda p: p.stat().st_mtime, reverse=True)
        return files[0]
    files.sort()
    return files[0]


def _agg(values: List[float]) -> Dict[str, Optional[float]]:
    """Compute mean, std, median over a list of floats (ignoring None/NaN)."""
    valid = [v for v in values if v is not None]
    if not valid:
        return {"n": 0, "mean": None, "std": None, "median": None}
    return {
        "n": len(valid),
        "mean": statistics.mean(valid),
        "std": statistics.pstdev(valid) if len(valid) >= 2 else None,
        "median": statistics.median(valid),
    }


def load_config_entries(config_path: Path) -> List[Dict[str, Any]]:
    """Load list of records from config JSON."""
    with config_path.open("r", encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, list):
        return data
    for key in ("items", "data", "samples"):
        if key in data and isinstance(data[key], list):
            return data[key]
    raise ValueError(f"Config must be a list or contain items/data/samples list: {config_path}")


def compute_stats(
    config_path: Path,
    *,
    results_key: str = "sam3d_results_path",
    metrics_glob: str = "metrics*.json",
    prefer_latest: bool = True,
    unique_paths: bool = True,
) -> Dict[str, Any]:
    """
    Aggregate scene-level and object-level metrics from metrics.json in each results dir.

    Returns dict with:
      - scene_level: predicted (and refined if present): mean, std, median over per-scene values
      - object_level: predicted (and refined if present): per-scene summary aggregation + optional per-object flatten
    """
    entries = load_config_entries(config_path)

    # Per-scene values (one value per metrics.json)
    scene_pred_cd: List[float] = []
    scene_pred_iou: List[float] = []
    scene_ref_cd: List[float] = []
    scene_ref_iou: List[float] = []
    obj_pred_cd: List[float] = []
    obj_pred_iou: List[float] = []
    obj_ref_cd: List[float] = []
    obj_ref_iou: List[float] = []

    # Per-object flatten (all objects across all scenes)
    all_obj_pred_cd: List[float] = []
    all_obj_pred_iou: List[float] = []
    all_obj_ref_cd: List[float] = []
    all_obj_ref_iou: List[float] = []

    missing_path = 0
    missing_dir = 0
    missing_metrics = 0
    parse_errors = 0
    used_metrics_files = 0
    seen: set[str] = set()
    per_sample: List[Dict[str, Any]] = []
    difficulty_stats: Dict[str, Dict[str, int]] = {
        k: {"total_records": 0, "used_metrics_files": 0} for k in _DIFFICULTY_KEYS
    }
    difficulty_metric_lists: Dict[str, Dict[str, List[float]]] = {
        k: {
            "scene_pred_cd": [],
            "scene_pred_iou": [],
            "scene_ref_cd": [],
            "scene_ref_iou": [],
            "obj_pred_cd": [],
            "obj_pred_iou": [],
            "obj_ref_cd": [],
            "obj_ref_iou": [],
        }
        for k in _DIFFICULTY_KEYS
    }

    for rec in entries:
        if not isinstance(rec, dict):
            continue
        difficulty = _difficulty_from_rec(rec)
        p = rec.get(results_key)
        if not isinstance(p, str) or not p.strip():
            missing_path += 1
            continue
        if unique_paths and p in seen:
            continue
        if unique_paths:
            seen.add(p)
        # Count total records per difficulty (after basic path validation)
        if difficulty not in difficulty_stats:
            difficulty = "unknown"
        difficulty_stats[difficulty]["total_records"] += 1

        d = Path(p)
        if not d.is_dir():
            missing_dir += 1
            continue

        mf = _pick_metrics_file(d, metrics_glob, prefer_latest=prefer_latest)
        if mf is None:
            missing_metrics += 1
            continue

        try:
            mjson = json.loads(mf.read_text(encoding="utf-8"))
        except Exception:
            parse_errors += 1
            continue
        used_metrics_files += 1
        difficulty_stats[difficulty]["used_metrics_files"] += 1

        # Scene-level predicted
        scene_cd_val = None
        scene_iou_val = None
        try:
            sl_pred = mjson.get("scene_level", {}).get("predicted", {})
            v = _safe_float(sl_pred.get("mean_cd"))
            if v is not None:
                scene_pred_cd.append(v)
                scene_cd_val = v
            v = _safe_float(sl_pred.get("mean_iou"))
            if v is not None:
                scene_pred_iou.append(v)
                scene_iou_val = v
        except Exception:
            pass

        # Scene-level refined
        scene_ref_cd_val = None
        scene_ref_iou_val = None
        try:
            sl_ref = mjson.get("scene_level", {}).get("refined", {})
            if sl_ref:
                v = _safe_float(sl_ref.get("mean_cd"))
                if v is not None:
                    scene_ref_cd.append(v)
                    scene_ref_cd_val = v
                v = _safe_float(sl_ref.get("mean_iou"))
                if v is not None:
                    scene_ref_iou.append(v)
                    scene_ref_iou_val = v
        except Exception:
            pass

        # Object-level predicted (per-scene summary)
        obj_cd_val = None
        obj_iou_val = None
        try:
            ol_pred = mjson.get("object_level_predicted", {}).get("summary", {})
            v = _safe_float(ol_pred.get("mean_cd"))
            if v is not None:
                obj_pred_cd.append(v)
                obj_cd_val = v
            v = _safe_float(ol_pred.get("mean_iou"))
            if v is not None:
                obj_pred_iou.append(v)
                obj_iou_val = v
        except Exception:
            pass

        # Object-level refined (per-scene summary)
        obj_ref_cd_val = None
        obj_ref_iou_val = None
        try:
            ol_ref = mjson.get("object_level_refined", {}).get("summary", {})
            if ol_ref:
                v = _safe_float(ol_ref.get("mean_cd"))
                if v is not None:
                    obj_ref_cd.append(v)
                    obj_ref_cd_val = v
                v = _safe_float(ol_ref.get("mean_iou"))
                if v is not None:
                    obj_ref_iou.append(v)
                    obj_ref_iou_val = v
        except Exception:
            pass

        metric_lists = difficulty_metric_lists.get(difficulty)
        if metric_lists is not None:
            if scene_cd_val is not None:
                metric_lists["scene_pred_cd"].append(scene_cd_val)
            if scene_iou_val is not None:
                metric_lists["scene_pred_iou"].append(scene_iou_val)
            if scene_ref_cd_val is not None:
                metric_lists["scene_ref_cd"].append(scene_ref_cd_val)
            if scene_ref_iou_val is not None:
                metric_lists["scene_ref_iou"].append(scene_ref_iou_val)
            if obj_cd_val is not None:
                metric_lists["obj_pred_cd"].append(obj_cd_val)
            if obj_iou_val is not None:
                metric_lists["obj_pred_iou"].append(obj_iou_val)
            if obj_ref_cd_val is not None:
                metric_lists["obj_ref_cd"].append(obj_ref_cd_val)
            if obj_ref_iou_val is not None:
                metric_lists["obj_ref_iou"].append(obj_ref_iou_val)

        # Per-sample row for Excel
        sample_id = p.strip()
        scene_id = rec.get("scene_id")
        image_id = rec.get("image_id")
        sub_scene_id = _sub_scene_id(rec)
        if scene_id is not None and image_id is not None:
            sample_id = f"{scene_id}_{image_id}"
        elif isinstance(rec.get("id"), (int, float)):
            sample_id = str(rec.get("id"))
        per_sample.append({
            "scene_id": scene_id,
            "image_id": image_id,
            "sub_scene_id": sub_scene_id,
            "sample_id": sample_id,
            "results_path": p.strip(),
            "scene_cd": scene_cd_val,
            "scene_iou": scene_iou_val,
            "object_cd": obj_cd_val,
            "object_iou": obj_iou_val,
            "scene_ref_cd": scene_ref_cd_val,
            "scene_ref_iou": scene_ref_iou_val,
            "object_ref_cd": obj_ref_cd_val,
            "object_ref_iou": obj_ref_iou_val,
        })

        # Per-object flatten: object_level_*["objects"]
        try:
            for obj in mjson.get("object_level_predicted", {}).get("objects", []):
                v = _safe_float(obj.get("chamfer_distance"))
                if v is not None:
                    all_obj_pred_cd.append(v)
                v = _safe_float(obj.get("iou"))
                if v is not None:
                    all_obj_pred_iou.append(v)
        except Exception:
            pass
        try:
            for obj in mjson.get("object_level_refined", {}).get("objects", []):
                v = _safe_float(obj.get("chamfer_distance"))
                if v is not None:
                    all_obj_ref_cd.append(v)
                v = _safe_float(obj.get("iou"))
                if v is not None:
                    all_obj_ref_iou.append(v)
        except Exception:
            pass

    for diff_stats in difficulty_stats.values():
        total = diff_stats.get("total_records", 0)
        used = diff_stats.get("used_metrics_files", 0)
        diff_stats["success_rate"] = used / total if total else None

    def build_level(name: str, cd_list: List[float], iou_list: List[float]) -> Dict[str, Any]:
        return {
            "n_scenes": len(cd_list),
            "mean_cd": _agg(cd_list),
            "mean_iou": _agg(iou_list),
        }

    def build_object_flatten(name: str, cd_list: List[float], iou_list: List[float]) -> Dict[str, Any]:
        return {
            "n_objects": len(cd_list),
            "mean_cd": _agg(cd_list),
            "mean_iou": _agg(iou_list),
        }

    difficulty_metrics: Dict[str, Any] = {}
    for difficulty, metric_lists in difficulty_metric_lists.items():
        scene_pred = build_level("predicted", metric_lists["scene_pred_cd"], metric_lists["scene_pred_iou"])
        scene_ref = None
        if metric_lists["scene_ref_cd"] or metric_lists["scene_ref_iou"]:
            scene_ref = build_level("refined", metric_lists["scene_ref_cd"], metric_lists["scene_ref_iou"])
        obj_pred = build_level("predicted", metric_lists["obj_pred_cd"], metric_lists["obj_pred_iou"])
        obj_ref = None
        if metric_lists["obj_ref_cd"] or metric_lists["obj_ref_iou"]:
            obj_ref = build_level("refined", metric_lists["obj_ref_cd"], metric_lists["obj_ref_iou"])
        difficulty_metrics[difficulty] = {
            "scene_level": {"predicted": scene_pred, "refined": scene_ref},
            "object_level": {"per_scene": {"predicted": obj_pred, "refined": obj_ref}},
        }

    return {
        "config": str(config_path),
        "results_key": results_key,
        "metrics_glob": metrics_glob,
        "total_records": len(entries),
        "unique_results_paths": len(seen),
        "used_metrics_files": used_metrics_files,
        "missing_results_path": missing_path,
        "missing_dir": missing_dir,
        "missing_metrics": missing_metrics,
        "parse_errors": parse_errors,
        "difficulty_stats": difficulty_stats,
        "difficulty_metrics": difficulty_metrics,
        "scene_level": {
            "predicted": build_level("predicted", scene_pred_cd, scene_pred_iou),
            "refined": build_level("refined", scene_ref_cd, scene_ref_iou) if (scene_ref_cd or scene_ref_iou) else None,
        },
        "object_level": {
            "per_scene": {
                "predicted": build_level("predicted", obj_pred_cd, obj_pred_iou),
                "refined": build_level("refined", obj_ref_cd, obj_ref_iou) if (obj_ref_cd or obj_ref_iou) else None,
            },
            "per_object": {
                "predicted": build_object_flatten("predicted", all_obj_pred_cd, all_obj_pred_iou),
                "refined": build_object_flatten("refined", all_obj_ref_cd, all_obj_ref_iou) if (all_obj_ref_cd or all_obj_ref_iou) else None,
            },
        },
        "per_sample": per_sample,
    }


def print_summary(stats: Dict[str, Any]) -> None:
    """Print human-readable summary to stdout."""
    print("=" * 60)
    print("SAM3D metrics statistics")
    print("=" * 60)
    print(f"Config: {stats['config']}")
    print(f"Results key: {stats['results_key']}")
    print(
        f"Unique results paths: {stats.get('unique_results_paths', 0)} (total records: {stats['total_records']}), "
        f"Parsed metrics files: {stats.get('used_metrics_files', 0)}, "
        f"Per-sample rows for Excel: {len(stats.get('per_sample') or [])}"
    )
    print(f"Missing path: {stats['missing_results_path']}, missing dir: {stats['missing_dir']}, missing metrics: {stats['missing_metrics']}, parse errors: {stats['parse_errors']}")
    print()

    difficulty_stats = stats.get("difficulty_stats") or {}
    if difficulty_stats:
        print("Difficulty success rates (parsed metrics / total records)")
        print("-" * 40)
        for key in _DIFFICULTY_KEYS:
            data = difficulty_stats.get(key)
            if not data:
                continue
            total = data.get("total_records", 0)
            used = data.get("used_metrics_files", 0)
            rate = data.get("success_rate")
            rate_s = f"{rate:.3f}" if isinstance(rate, (int, float)) else "N/A"
            print(f"  {key}: {used}/{total} ({rate_s})")
        print()

    def _print_agg(label: str, agg: Dict[str, Any]) -> None:
        if agg.get("n", 0) == 0:
            print(f"  {label}: (no data)")
            return
        m = agg.get("mean")
        s = agg.get("std")
        med = agg.get("median")
        n = agg.get("n", 0)
        mean_s = f"{m:.6f}" if m is not None else "N/A"
        std_s = f"{s:.6f}" if s is not None else "N/A"
        med_s = f"{med:.6f}" if med is not None else "N/A"
        print(f"  {label}: n={n}  mean={mean_s}  std={std_s}  median={med_s}")

    print("Scene-level (per-scene mean CD / mean IoU)")
    print("-" * 40)
    sl = stats.get("scene_level", {})
    for key in ("predicted", "refined"):
        lev = sl.get(key)
        if lev is None:
            continue
        print(f"  [{key}]")
        _print_agg("  mean_cd", lev["mean_cd"])
        _print_agg("  mean_iou", lev["mean_iou"])
    print()

    print("Object-level per-scene (one value per scene from summary)")
    print("-" * 40)
    ol = stats.get("object_level", {}).get("per_scene", {})
    for key in ("predicted", "refined"):
        lev = ol.get(key)
        if lev is None:
            continue
        print(f"  [{key}]")
        _print_agg("  mean_cd", lev["mean_cd"])
        _print_agg("  mean_iou", lev["mean_iou"])
    print()

    print("Object-level per-object (all objects flattened)")
    print("-" * 40)
    ol_flat = stats.get("object_level", {}).get("per_object", {})
    for key in ("predicted", "refined"):
        lev = ol_flat.get(key)
        if lev is None:
            continue
        print(f"  [{key}]")
        _print_agg("  mean_cd", lev["mean_cd"])
        _print_agg("  mean_iou", lev["mean_iou"])
    print("=" * 60)


def _agg_row(level: str, variant: str, metric: str, agg: Dict[str, Any]) -> Dict[str, Any]:
    """One row for Excel: level, variant, metric, n, mean, std, median."""
    n = agg.get("n") or 0
    mean = agg.get("mean")
    std = agg.get("std")
    median = agg.get("median")
    return {
        "level": level,
        "variant": variant,
        "metric": metric,
        "n": n,
        "mean": round(mean, 6) if mean is not None and mean == mean else None,
        "std": round(std, 6) if std is not None and std == std else None,
        "median": round(median, 6) if median is not None and median == median else None,
    }


def stats_to_excel_rows(stats: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Flatten stats dict into rows for Excel (level, variant, metric, n, mean, std, median)."""
    rows: List[Dict[str, Any]] = []
    # Meta
    rows.append({"level": "meta", "variant": "config", "metric": stats.get("config", ""), "n": "", "mean": "", "std": "", "median": ""})
    rows.append({"level": "meta", "variant": "results_key", "metric": stats.get("results_key", ""), "n": "", "mean": "", "std": "", "median": ""})
    rows.append({"level": "meta", "variant": "unique_results_paths", "metric": str(stats.get("unique_results_paths", 0)), "n": stats.get("total_records", 0), "mean": "", "std": "", "median": ""})
    rows.append({"level": "meta", "variant": "parsed_metrics_files", "metric": str(stats.get("used_metrics_files", 0)), "n": "", "mean": "", "std": "", "median": ""})
    rows.append({"level": "meta", "variant": "missing", "metric": f"path={stats.get('missing_results_path',0)} dir={stats.get('missing_dir',0)} metrics={stats.get('missing_metrics',0)} parse={stats.get('parse_errors',0)}", "n": "", "mean": "", "std": "", "median": ""})
    # Scene-level
    for key in ("predicted", "refined"):
        lev = stats.get("scene_level", {}).get(key)
        if lev is None:
            continue
        for m in ("mean_cd", "mean_iou"):
            agg = lev.get(m, {})
            rows.append(_agg_row("scene_level", key, m, agg))
    # Object-level per_scene
    for key in ("predicted", "refined"):
        lev = stats.get("object_level", {}).get("per_scene", {}).get(key)
        if lev is None:
            continue
        for m in ("mean_cd", "mean_iou"):
            agg = lev.get(m, {})
            rows.append(_agg_row("object_per_scene", key, m, agg))
    # Object-level per_object
    for key in ("predicted", "refined"):
        lev = stats.get("object_level", {}).get("per_object", {}).get(key)
        if lev is None:
            continue
        for m in ("mean_cd", "mean_iou"):
            agg = lev.get(m, {})
            rows.append(_agg_row("object_per_object", key, m, agg))
    return rows


def _extract_per_sample_map(stats: Dict[str, Any], prefix: str) -> Dict[str, Dict[str, Any]]:
    out: Dict[str, Dict[str, Any]] = {}
    per_sample = stats.get("per_sample") or []
    for row in per_sample:
        sid = row.get("sample_id") or row.get("results_path")
        if sid is None:
            continue
        sid = str(sid)
        dest = out.setdefault(sid, {})
        for key in ("scene_id", "image_id", "sub_scene_id"):
            val = row.get(key)
            if val is not None and val != "":
                dest.setdefault(key, val)
        for metric in _METRIC_KEYS:
            dest[f"{prefix}_{metric}"] = row.get(metric)
    return out


def _merge_sample_maps(maps: List[Dict[str, Dict[str, Any]]]) -> Dict[str, Dict[str, Any]]:
    combined: Dict[str, Dict[str, Any]] = {}
    for m in maps:
        for sid, vals in m.items():
            dst = combined.setdefault(sid, {})
            dst.update(vals)
    return combined


def build_combined_per_sample(
    config_path: Path,
    *,
    results_keys: List[str],
    metrics_glob: str,
    prefer_latest: bool,
    unique_paths: bool,
) -> Dict[str, Any]:
    """Build combined per-sample metrics for multiple results keys."""
    maps: List[Dict[str, Dict[str, Any]]] = []
    prefixes: List[str] = []
    for results_key in results_keys:
        prefix = _prefix_for_results_key(results_key)
        prefixes.append(prefix)
        stats = compute_stats(
            config_path,
            results_key=results_key,
            metrics_glob=_combined_metrics_glob(results_key, metrics_glob),
            prefer_latest=prefer_latest,
            unique_paths=unique_paths,
        )
        maps.append(_extract_per_sample_map(stats, prefix))

    combined = _merge_sample_maps(maps)

    def sort_key(item: tuple[str, Dict[str, Any]]) -> tuple[str, str, str, str]:
        data = item[1]
        return (
            str(data.get("scene_id", "")),
            str(data.get("image_id", "")),
            str(data.get("sub_scene_id", "")),
            item[0],
        )

    rows: List[Dict[str, Any]] = []
    for _, data in sorted(combined.items(), key=sort_key):
        row: Dict[str, Any] = {
            "scene_id": data.get("scene_id"),
            "image_id": data.get("image_id"),
            "sub_scene_id": data.get("sub_scene_id"),
        }
        for prefix in prefixes:
            for metric in _METRIC_KEYS:
                row[f"{prefix}_{metric}"] = data.get(f"{prefix}_{metric}")
        rows.append(row)
    return {"prefixes": prefixes, "rows": rows}


def write_stats_to_excel(
    stats: Dict[str, Any],
    output_path: Path,
    *,
    combined_per_sample: Optional[Dict[str, Any]] = None,
) -> None:
    """Write stats to Excel: per-sample + combined + aggregate sheets."""
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        raise ImportError("Writing Excel requires openpyxl. Install with: pip install openpyxl") from None

    wb = openpyxl.Workbook()
    per_sample = stats.get("per_sample") or []

    if combined_per_sample:
        combined_rows = combined_per_sample.get("rows") or []
        prefixes = combined_per_sample.get("prefixes") or []
        ws_combined = wb.create_sheet("Combined_per_sample", 0)
        combined_headers = ["scene_id", "image_id", "sub_scene_id"]
        for prefix in prefixes:
            for metric in _METRIC_KEYS:
                combined_headers.append(f"{prefix}_{metric}")
        for c, h in enumerate(combined_headers, 1):
            ws_combined.cell(row=1, column=c, value=h)
        ws_combined.row_dimensions[1].font = Font(bold=True)
        for r, row in enumerate(combined_rows, 2):
            for c, key in enumerate(combined_headers, 1):
                val = row.get(key)
                if isinstance(val, float):
                    val = round(val, 6)
                ws_combined.cell(row=r, column=c, value=val)

    # Sheet 1 (default): per-sample CD/IoU rows.
    ws1 = wb.active
    ws1.title = "Per_sample"
    sample_headers = [
        "scene_id", "image_id", "sub_scene_id",
        "sample_id", "results_path",
        "scene_cd", "scene_iou", "object_cd", "object_iou",
    ]
    for c, h in enumerate(sample_headers, 1):
        ws1.cell(row=1, column=c, value=h)
    ws1.row_dimensions[1].font = Font(bold=True)
    for r, row in enumerate(per_sample, 2):
        for c, key in enumerate(sample_headers, 1):
            val = row.get(key)
            if isinstance(val, float):
                val = round(val, 6)
            ws1.cell(row=r, column=c, value=val)

    # Sheet 2: aggregate statistics.
    ws2 = wb.create_sheet("Stats", 1)
    headers = ["level", "variant", "metric", "n", "mean", "std", "median"]
    for c, h in enumerate(headers, 1):
        ws2.cell(row=1, column=c, value=h)
    ws2.row_dimensions[1].font = Font(bold=True)
    rows = stats_to_excel_rows(stats)
    for r, row in enumerate(rows, 2):
        for c, key in enumerate(headers, 1):
            val = row.get(key)
            ws2.cell(row=r, column=c, value=val)

    wb.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Compute object-level and scene-level statistics from metrics.json under sam3d_results_path."
    )
    parser.add_argument(
        "--config",
        type=Path,
        required=True,
        help="Path to config JSON (list of records with sam3d_results_path).",
    )
    parser.add_argument(
        "--results-key",
        type=str,
        default="sam3d_results_path",
        help="Key in each record for results directory (default: sam3d_results_path).",
    )
    parser.add_argument(
        "--metrics-glob",
        type=str,
        default="metrics*.json",
        help="Glob for metrics file in each results dir (default: metrics*.json).",
    )
    parser.add_argument("--prefer-latest", action="store_true", help="If multiple metrics files, use latest by mtime.")
    parser.add_argument("--no-unique-paths", action="store_true", help="Do not deduplicate by results path.")
    parser.add_argument(
        "--combined-results-keys",
        type=str,
        default=",".join(_DEFAULT_COMBINED_RESULTS_KEYS),
        help=(
            "Comma-separated results keys for combined per-sample Excel sheet "
            "(empty to disable)."
        ),
    )
    parser.add_argument(
        "--out-json",
        type=Path,
        default=None,
        help="Write full stats to this JSON file.",
    )
    parser.add_argument(
        "--out-excel",
        type=Path,
        default=None,
        help="Write stats to this Excel file (requires openpyxl).",
    )
    parser.add_argument("--quiet", action="store_true", help="Do not print summary to stdout.")
    args = parser.parse_args()

    stats = compute_stats(
        args.config,
        results_key=args.results_key,
        metrics_glob=args.metrics_glob,
        prefer_latest=args.prefer_latest,
        unique_paths=not args.no_unique_paths,
    )

    if not args.quiet:
        print_summary(stats)

    if args.out_json is not None:
        args.out_json.parent.mkdir(parents=True, exist_ok=True)
        # Serialize for JSON (nan -> null)
        def _to_serializable(obj: Any) -> Any:
            if isinstance(obj, dict):
                return {k: _to_serializable(v) for k, v in obj.items()}
            if isinstance(obj, list):
                return [_to_serializable(v) for v in obj]
            if isinstance(obj, float) and (obj != obj or obj == float("inf") or obj == float("-inf")):
                return None
            return obj

        with args.out_json.open("w", encoding="utf-8") as f:
            json.dump(_to_serializable(stats), f, indent=2)
        print(f"Wrote: {args.out_json}")

    if args.out_excel is not None:
        out_xlsx = args.out_excel
        if out_xlsx.exists() and out_xlsx.is_dir():
            out_xlsx = out_xlsx / f"sam3d_metrics_stats_{args.config.stem}_{args.results_key}.xlsx"
        elif str(out_xlsx).endswith("/") or str(out_xlsx).endswith("\\"):
            out_xlsx = out_xlsx / f"sam3d_metrics_stats_{args.config.stem}_{args.results_key}.xlsx"
        elif out_xlsx.suffix.lower() != ".xlsx":
            out_xlsx = out_xlsx.with_suffix(".xlsx")
        out_xlsx.parent.mkdir(parents=True, exist_ok=True)
        combined_keys = [k.strip() for k in (args.combined_results_keys or "").split(",") if k.strip()]
        combined_per_sample = None
        if combined_keys:
            combined_per_sample = build_combined_per_sample(
                args.config,
                results_keys=combined_keys,
                metrics_glob=args.metrics_glob,
                prefer_latest=args.prefer_latest,
                unique_paths=not args.no_unique_paths,
            )
        write_stats_to_excel(stats, out_xlsx, combined_per_sample=combined_per_sample)
        print(f"Wrote: {out_xlsx}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
